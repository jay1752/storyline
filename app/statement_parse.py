"""Reads a broker export (CSV / TSV / XLSX / XLS) into statement rows without an AI model.

Broker files differ in column names but not in meaning, so instead of one parser per broker this finds the header row, maps each column by
what it holds (ISIN, symbol, name, quantity, average price, trade price, buy/sell, date) and builds the same Txn rows the AI path produces,
which then go through the same matching, price check and review card (app/positions.py). Exact where it recognises the file, it never
guesses: anything it cannot map confidently returns None and the chat falls back to the AI reading the file's text.

Two shapes are recognised:
  * holdings   - one row per stock with a quantity (and usually an average cost); no buy/sell column.
  * trades     - one row per trade with a buy/sell column, quantity, price and date.
A profit-and-loss / realised-gains report looks like holdings but lists what was SOLD; it is refused with a plain message instead of
being turned into positions."""
import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from app.positions import Txn

MAX_ROWS = 20000


@dataclass
class Parsed:
    kind: str                       # holdings | trades
    txns: list
    mapping: str                    # human-readable "how I read it"
    skipped: int = 0
    notes: list = field(default_factory=list)
    market: str | None = None      # "IN" when the file is plainly Indian (NSE/BSE exchange column or INE/INF ISINs)


class NotPositions(Exception):
    """The file is recognisable but is not holdings or trades (e.g. a realised P&L report)."""


# ---------------------------------------------------------------------------------------------- reading files

def read_sheets(filename: str, data: bytes) -> list[tuple[str, list[list]]]:
    """[(sheet name, rows)] with cells as str / number / datetime. Raises ValueError for an unreadable file."""
    name = filename.lower()
    if name.endswith(".pdf"):
        from app.pdf_tables import pdf_sheets
        sheets = pdf_sheets(data)
        if not sheets:
            raise ValueError("no readable text layer in this PDF")
        return sheets
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_ROWS:
                    break
                rows.append(list(row))
            out.append((ws.title, rows))
        return out
    if name.endswith(".xls"):
        import xlrd
        wb = xlrd.open_workbook(file_contents=data)
        out = []
        for sh in wb.sheets():
            rows = []
            for r in range(min(sh.nrows, MAX_ROWS)):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            row.append(datetime(*xlrd.xldate_as_tuple(cell.value, wb.datemode)))
                            continue
                        except Exception:  # noqa: BLE001
                            pass
                    row.append(cell.value if cell.value != "" else None)
                rows.append(row)
            out.append((sh.name, rows))
        return out
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    sample = text[:4096]
    delim = "\t" if sample.count("\t") > sample.count(",") else (";" if sample.count(";") > sample.count(",") else ",")
    rows = [r for r in csv.reader(io.StringIO(text), delimiter=delim)][:MAX_ROWS]
    return [("file", rows)]


def sheets_to_text(sheets: list[tuple[str, list[list]]], limit: int = 400_000) -> str:
    """Plain CSV-ish text of every sheet, for the AI path when the columns could not be mapped."""
    parts, size = [], 0
    for name, rows in sheets:
        parts.append(f"## sheet: {name}")
        for r in rows:
            line = ",".join("" if c is None else (c.date().isoformat() if isinstance(c, datetime) else str(c)).replace("\n", " ") for c in r).rstrip(",")
            if not line:
                continue
            size += len(line) + 1
            if size > limit:
                return "\n".join(parts) + "\n[file truncated]"
            parts.append(line)
    return "\n".join(parts)


# ---------------------------------------------------------------------------------------------- column mapping

def _norm(h) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("₹", "").replace("$", "").replace(",", "").replace(" ", "")
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if s in ("", "-", "--", "NA", "N/A", "nan"):
        return None
    try:
        x = float(s)
    except ValueError:
        return None
    return -x if neg else x


_DATE_FORMATS = ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d %b %Y", "%d-%b-%y", "%d/%m/%y", "%Y/%m/%d", "%d-%B-%Y", "%m/%d/%Y", "%b %d %Y", "%B %d %Y", "%b %d, %Y", "%d.%m.%Y", "%d %B %Y")


def _date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()[:25]
    for cut in (s, s[:10], s.split(" ")[0], s.split("T")[0]):
        for f in _DATE_FORMATS:
            try:
                return datetime.strptime(cut, f).date().isoformat()
            except ValueError:
                continue
    return None


def _first(norms, *tests, exclude=()):
    """Index of the first header (in priority order of the tests) that equals / contains a test and none of `exclude`."""
    for t in tests:
        for i, h in enumerate(norms):
            if not h or any(x in h for x in exclude):
                continue
            if h == t or (len(t) > 4 and t in h):
                return i
    return None


def _map_columns(header: list) -> dict | None:
    n = [_norm(h) for h in header]
    isin = next((i for i, h in enumerate(n) if "isin" in h and "name" not in h and "desc" not in h), None)
    sym = _first(n, "symbol", "ticker", "tradingsymbol", "scripsymbol", "scripcode", "instrument", "scrip", "securitysymbol")
    name = _first(n, "companyname", "stockname", "securityname", "scripname", "isinname", "instrumentname", "company", "security", "stock", "name", exclude=("client", "exchange", "broker", "customer", "account"))
    side = next((i for tst in ("tradetype", "buysell", "buyorsell", "transactiontype", "trantype", "side", "action", "orderside", "type", "buysellflag", "tradeside")
                 for i, h in enumerate(n) if h == tst), None)
    qty = _first(n, "totqty", "totalqty", "totalquantity", "quantityheld", "qtyheld", "quantityavailable", "quantity", "qty", "tradedquantity", "tradedqty", "netqty", "holdingquantity", "totalquantity", "shares", "units", "currentbalance", "balance",
                 exclude=("pledge", "longterm", "discrepant", "value", "amount", "price", "rate", "cost", "pnl", "unrealised", "unrealized", "t1", "collateral", "free"))
    price_hold = _first(n, "averageprice", "avgprice", "averagecost", "avgcost", "averagebuyprice", "avgbuyprice", "buyaverage", "buyavg", "costprice", "purchaseprice", "averagecostprice", "avgcostprice", "acquisitionprice",
                        exclude=("total", "value", "previous", "closing", "ltp", "market", "current"))
    price_trade = _first(n, "tradeprice", "netrate", "executionprice", "tradedprice", "rate", "price", "avgprice", "averageprice", exclude=("previous", "closing", "ltp", "market", "current", "total", "value", "buy", "sell", "stop", "trigger"))
    dt = _first(n, "tradedate", "transactiondate", "executiondate", "orderexecutiontime", "purchasedate", "buydate", "datebought", "dateofpurchase", "acquisitiondate", "datepurchased", "boughton", "date", "tradetime", exclude=("expiry", "settlement", "maturity"))
    if dt is None:  # a header wrapped over two lines in a PDF can leave just "Transaction" / "Trade"
        dt = _first(n, "transaction", "trade", "executed", "orderdate")
    if qty is None or not (isin is not None or sym is not None or name is not None):
        return None
    remarks = _first(n, "remarks", "narration", "description", "particulars")
    extra_qty = [i for i, h in enumerate(n) if i != qty and h and ("pledge" in h or "t1" in h or "collateral" in h) and ("qty" in h or "quantity" in h or "units" in h) and "value" not in h and "amount" not in h]
    seg = _first(n, "segment", "exchangesegment")
    exch = next((i for i, h in enumerate(n) if h in ("exchange", "exch", "market")), None)
    return {"isin": isin, "sym": sym, "name": name, "side": side, "qty": qty, "extra_qty": extra_qty, "price_hold": price_hold, "price_trade": price_trade, "date": dt, "seg": seg, "remarks": remarks, "exch": exch,
            "sellmark": any(re.search(r"sellvalue|sellprice|sellamount|(?<!un)realised|(?<!un)realized|profitloss", h) for h in n)}


_SIDE = {"b": "buy", "buy": "buy", "purchase": "buy", "bought": "buy", "s": "sell", "sell": "sell", "sale": "sell", "sold": "sell"}
_SKIP_WORDS = ("total", "grand total", "net total", "sub total", "subtotal")


def _clean_symbol(s: str) -> str:
    s = str(s).strip().upper()
    return re.sub(r"-(EQ|BE|BZ|SM|ST|N1|N2)$", "", s)


def _parse_sheet(rows: list[list]) -> Parsed | None:
    header_at, cols = None, None
    for i, r in enumerate(rows[:60]):
        m = _map_columns(r)
        if m and sum(1 for c in r if c not in (None, "")) >= 3:
            header_at, cols = i, m
            break
    if cols is None:
        return None
    is_trades = cols["side"] is not None
    if not is_trades and cols["sellmark"] and cols["price_hold"] is None:
        raise NotPositions("This looks like a profit-and-loss / realised-gains report: it lists what you SOLD, not what you hold. "
                           "Upload your holdings statement or tradebook instead.")
    txns, skipped, corp = [], 0, 0
    exchanges: set[str] = set()
    for r in rows[header_at + 1:]:
        if not r or all(c in (None, "") for c in r):
            continue

        def cell(k):
            i = cols.get(k)
            return r[i] if i is not None and i < len(r) else None

        isin = str(cell("isin") or "").strip().upper()
        isin = isin if re.fullmatch(r"[A-Z]{2}[A-Z0-9]{9}[0-9]", isin) else None
        sym, nm = cell("sym"), cell("name")
        label = (_clean_symbol(sym) if sym not in (None, "") else (str(nm).strip() if nm not in (None, "") else (isin or "")))
        if not label or any(label.lower().startswith(w) for w in _SKIP_WORDS):
            continue
        ex = str(cell("exch") or "").strip().upper()
        if ex:
            exchanges.add(ex)
        seg = str(cell("seg") or "").upper()
        if seg and any(k in seg for k in ("FO", "FUT", "OPT", "CD", "COM", "MCX")) and "EQ" not in seg:
            skipped += 1
            continue
        q = _num(cell("qty"))
        if q is None:
            continue
        for xi in cols["extra_qty"]:
            q += _num(r[xi] if xi < len(r) else None) or 0
        if is_trades:
            side = _SIDE.get(str(cell("side") or "").strip().lower().rstrip("."))
            if side is None:
                skipped += 1
                continue
            price = _num(cell("price_trade"))
            if q < 0:
                q = -q
            if q == 0:
                continue
            if price == 0 and re.search(r"split|bonus", str(cell("remarks") or ""), re.I):
                corp += 1
            txns.append(Txn(label, side, q, price, _date(cell("date")), None, isin, None))
        else:
            if q <= 0:
                skipped += 1
                continue
            price = _num(cell("price_hold"))
            txns.append(Txn(label, "holding", q, price if price and price > 0 else None, _date(cell("date")), None, isin, None))
    if len(txns) < 1:
        return None
    hdr = rows[header_at]
    col = lambda k: str(hdr[cols[k]]).strip() if cols.get(k) is not None else None
    parts = [f"{lab} = “{col(k)}”" for lab, k in (("company", "sym" if cols["sym"] is not None else "name"), ("ISIN", "isin"), ("quantity", "qty"),
                                                   ("price", "price_trade" if is_trades else "price_hold"), ("buy/sell", "side"), ("date", "date")) if col(k)]
    if cols["extra_qty"]:
        parts.append("+ " + ", ".join(f"“{hdr[i]}”" for i in cols["extra_qty"]))
    notes = []
    if corp:
        notes.append(f"{corp} split/bonus row{'s' if corp != 1 else ''} counted as free shares, so average cost is adjusted for them.")
    if not is_trades and all(t.price is None for t in txns):
        notes.append("This file has no average price, so cost has to be added by hand.")
    india = (exchanges and exchanges <= {"NSE", "BSE", "NSEI", "BSEI", "NSE_EQ", "BSE_EQ"}) or (txns and all((t.isin or "").startswith("IN") for t in txns))
    return Parsed("trades" if is_trades else "holdings", txns, "; ".join(parts), skipped, notes, "IN" if india else None)


def parse_statement(sheets: list[tuple[str, list[list]]]) -> Parsed | None:
    """The first sheet that reads as holdings or trades. Raises NotPositions for a recognised non-positions report. None = unknown layout."""
    refusal = None
    for _, rows in sheets:
        try:
            p = _parse_sheet(rows)
        except NotPositions as e:
            refusal = e
            continue
        if p:
            return p
    if refusal:
        raise refusal
    return None
